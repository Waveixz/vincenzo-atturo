import re
from datetime import date, datetime
from io import BytesIO
from pathlib import Path

import streamlit as st
from docx import Document

from core.storage import DATA_DIR, add_record, delete_record, load_data, safe_filename, save_upload, update_record
from core.ui import empty_state, module_header, rerun_notice


KIT_PATH = Path(r"C:\Users\vatturo\Desktop\vincenzo-atturo\output\commercial-kit")
GENERATED_DIR = DATA_DIR / "generated"
PLACEHOLDER_PATTERN = re.compile(r"\[[^\[\]\r\n]{1,100}\]")


def _display_name(path: Path) -> str:
    name = re.sub(r"^\d+_", "", path.stem).replace("_", " ")
    return name.capitalize()


def _iter_paragraphs(document):
    for paragraph in document.paragraphs:
        yield paragraph
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from _iter_paragraphs(cell)
    for section in getattr(document, "sections", []):
        for area in (section.header, section.footer):
            for paragraph in area.paragraphs:
                yield paragraph
            for table in area.tables:
                for row in table.rows:
                    for cell in row.cells:
                        yield from _iter_paragraphs(cell)


def _extract_placeholders(path: Path) -> list[str]:
    document = Document(path)
    found = set()
    for paragraph in _iter_paragraphs(document):
        found.update(PLACEHOLDER_PATTERN.findall(paragraph.text))
    return sorted(found, key=str.casefold)


def _replace_paragraph(paragraph, replacements: dict[str, str]) -> None:
    for run in paragraph.runs:
        for token, value in replacements.items():
            if token in run.text:
                run.text = run.text.replace(token, value)

    unresolved = [token for token in replacements if token in paragraph.text]
    if not unresolved:
        return

    full_text = paragraph.text
    for token in unresolved:
        full_text = full_text.replace(token, replacements[token])
    if paragraph.runs:
        paragraph.runs[0].text = full_text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(full_text)


def _build_document(template: Path, replacements: dict[str, str]) -> bytes:
    document = Document(template)
    clean_replacements = {key: value.strip() for key, value in replacements.items() if value.strip()}
    for paragraph in _iter_paragraphs(document):
        _replace_paragraph(paragraph, clean_replacements)
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _template_files() -> list[Path]:
    if not KIT_PATH.exists():
        return []
    return sorted(KIT_PATH.glob("*.docx"))


def _common_replacements(values: dict[str, str]) -> dict[str, str]:
    aliases = {
        "client": ["[Cliente]", "[cliente]", "[denominazione del cliente]", "[Ragione sociale o nominativo]"],
        "project": ["[Progetto]", "[progetto]", "[Titolo]", "[Titolo del progetto]"],
        "date": ["[Data]", "[data]"],
        "contact": ["[Nome]", "[nome]", "[Nome e ruolo]", "[nome e ruolo]", "[Nome, ruolo e contatto]"],
        "email": ["[Email]", "[email]"],
        "phone": ["[Telefono]", "[telefono]"],
        "address": ["[Indirizzo]", "[indirizzo]"],
        "tax": ["[Dati fiscali]", "[dati fiscali]"],
        "amount": ["[Importo]", "[importo]"],
        "reference": ["[Riferimento]", "[Numero e data del preventivo]", "[VA-AAAA-NNN]", "[VAR-AAAA-NNN]"],
    }
    return {
        token: values[field]
        for field, tokens in aliases.items()
        for token in tokens
        if values.get(field, "").strip()
    }


def _render_archive():
    data = load_data()
    documents = data["documents"]
    col1, col2 = st.columns([2, 1])
    query = col1.text_input("Cerca", placeholder="Titolo, categoria, cliente o nota")
    category = col2.selectbox("Categoria", ["Tutte", "Contratti", "Preventivi", "Clienti", "Progetti", "Tecnici", "Altro"])
    filtered = documents
    if category != "Tutte":
        filtered = [item for item in filtered if item.get("category") == category]
    if query:
        needle = query.casefold()
        filtered = [item for item in filtered if needle in " ".join(str(value) for value in item.values()).casefold()]
    if filtered:
        st.dataframe([
            {"Titolo": item.get("title", ""), "Categoria": item.get("category", ""), "Cliente o progetto": item.get("context", ""), "File": Path(item.get("path", "")).name, "Note": item.get("notes", "")}
            for item in filtered
        ], hide_index=True, width="stretch")
        with st.expander("Gestisci documento · modifica / elimina", icon=":material/edit:"):
            options = {item["id"]: item["title"] for item in documents}
            selected_id = st.selectbox("Documento", options, format_func=options.get, key="document_manage_id")
            selected = next(item for item in documents if item["id"] == selected_id)
            c1, c2 = st.columns(2)
            title_e = c1.text_input("Titolo", value=selected.get("title", ""), key="document_edit_title")
            cats = ["Contratti", "Preventivi", "Clienti", "Progetti", "Tecnici", "Altro"]
            cat_e = c2.selectbox("Categoria", cats, index=cats.index(selected.get("category")) if selected.get("category") in cats else len(cats)-1, key="document_edit_category")
            context_e = c1.text_input("Cliente o progetto", value=selected.get("context", ""), key="document_edit_context")
            path_e = c2.text_input("Percorso o URL", value=selected.get("path", ""), key="document_edit_path")
            notes_e = st.text_area("Note", value=selected.get("notes", ""), key="document_edit_notes")
            if st.button("Salva modifiche", icon=":material/save:", key="document_edit_save"):
                if not title_e.strip(): st.error("Inserisci un titolo.")
                else:
                    update_record("documents", selected_id, {"title": title_e.strip(), "category": cat_e, "context": context_e.strip(), "path": path_e.strip(), "notes": notes_e.strip()})
                    rerun_notice("Documento aggiornato")
            confirm = st.checkbox("Confermo la rimozione del riferimento dall’archivio", key=f"document_confirm_{selected_id}")
            if st.button("Rimuovi definitivamente", icon=":material/delete:", disabled=not confirm, key="document_delete"):
                delete_record("documents", selected_id)
                rerun_notice("Riferimento rimosso")
    else:
        empty_state("Archivio vuoto", "Compila un modello, carica un documento o registra un percorso esistente.")


def _render_templates(templates: list[Path]):
    st.subheader("Modelli VA Digital")
    if not templates:
        st.warning("La cartella del kit commerciale non è disponibile in questa postazione.")
        return
    selected = st.selectbox("Documento", templates, format_func=_display_name, key="template_download")
    st.caption(f"Modello modificabile Word · {round(selected.stat().st_size / 1024, 1)} KB")
    st.download_button(
        "Scarica il modello originale",
        data=selected.read_bytes(),
        file_name=selected.name,
        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        icon=":material/download:",
    )
    st.dataframe(
        [{"Modello": _display_name(path), "Formato": "DOCX", "Dimensione KB": round(path.stat().st_size / 1024, 1)} for path in templates],
        hide_index=True,
        width="stretch",
    )


def _render_compiler(templates: list[Path]):
    st.subheader("Compila un documento")
    if not templates:
        st.warning("Nessun modello Word disponibile.")
        return

    selected = st.selectbox("Modello da compilare", templates, format_func=_display_name, key="template_compile")
    placeholders = _extract_placeholders(selected)
    sample = {key: "x" for key in ["client", "project", "date", "contact", "email", "phone", "address", "tax", "amount", "reference"]}
    common_tokens = set(_common_replacements(sample))
    specific_tokens = [
        token for token in placeholders
        if token not in common_tokens and token != "[ ]" and len(token) <= 70
    ]

    with st.form("compile_document_form"):
        left, right = st.columns(2)
        client = left.text_input("Cliente o azienda")
        project = right.text_input("Progetto o oggetto")
        contact = left.text_input("Referente e ruolo")
        document_date = right.date_input("Data", value=date.today(), format="DD/MM/YYYY")
        email = left.text_input("Email")
        phone = right.text_input("Telefono")
        address = left.text_input("Indirizzo")
        tax = right.text_input("Partita IVA o codice fiscale")
        amount = left.text_input("Importo", placeholder="es. 2.500,00 euro + IVA")
        reference = right.text_input("Numero o riferimento documento")

        custom_values = {}
        if specific_tokens:
            with st.expander("Campi specifici del modello"):
                st.caption("Compila solo i campi utili: quelli lasciati vuoti resteranno modificabili nel file Word.")
                columns = st.columns(2)
                for index, token in enumerate(specific_tokens[:30]):
                    label = token[1:-1]
                    custom_values[token] = columns[index % 2].text_input(label, key=f"field_{selected.stem}_{index}")
                if len(specific_tokens) > 30:
                    st.caption(f"Altri {len(specific_tokens) - 30} campi resteranno disponibili nel documento Word.")

        submitted = st.form_submit_button("Genera documento Word", icon=":material/description:")

    if submitted:
        values = {
            "client": client,
            "project": project,
            "date": document_date.strftime("%d/%m/%Y"),
            "contact": contact,
            "email": email,
            "phone": phone,
            "address": address,
            "tax": tax,
            "amount": amount,
            "reference": reference,
        }
        replacements = _common_replacements(values)
        replacements.update(custom_values)
        content = _build_document(selected, replacements)
        context = client.strip() or project.strip() or "documento"
        timestamp = datetime.now().strftime("%Y%m%d_%H%M")
        output_name = safe_filename(f"{timestamp}_{context}_{selected.name}")
        GENERATED_DIR.mkdir(parents=True, exist_ok=True)
        output_path = GENERATED_DIR / output_name
        output_path.write_bytes(content)
        add_record("documents", {
            "title": f"{_display_name(selected)} - {context}",
            "category": "Contratti" if "contratto" in selected.stem else "Progetti",
            "context": context,
            "path": str(output_path),
            "notes": f"Generato dal modello {selected.name}",
        })
        st.session_state["generated_document"] = content
        st.session_state["generated_document_name"] = output_name
        st.success("Documento generato e aggiunto all'archivio locale.")

    if st.session_state.get("generated_document"):
        st.download_button(
            "Scarica il documento compilato",
            data=st.session_state["generated_document"],
            file_name=st.session_state["generated_document_name"],
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            icon=":material/download:",
            type="primary",
        )


def _render_upload():
    mode = st.segmented_control("Origine", ["Carica file", "Registra percorso"], default="Carica file")
    with st.form("document_form", clear_on_submit=True):
        col1, col2 = st.columns(2)
        title = col1.text_input("Titolo")
        category = col2.selectbox("Categoria", ["Contratti", "Preventivi", "Clienti", "Progetti", "Tecnici", "Altro"])
        context = col1.text_input("Cliente o progetto")
        if mode == "Carica file":
            uploaded = col2.file_uploader("File", type=["pdf", "docx", "xlsx", "csv", "xml", "p7m", "zip", "txt", "png", "jpg", "jpeg"])
            path_value = ""
        else:
            uploaded = None
            path_value = col2.text_input("Percorso completo o URL")
        notes = st.text_area("Note")
        if st.form_submit_button("Salva in archivio", icon=":material/save:"):
            if not title.strip():
                st.error("Inserisci un titolo.")
            elif mode == "Carica file" and uploaded is None:
                st.error("Seleziona un file.")
            elif mode == "Registra percorso" and not path_value.strip():
                st.error("Inserisci un percorso o un URL.")
            else:
                if uploaded is not None:
                    saved_path = save_upload(uploaded.name, uploaded.getvalue())
                    path_value = str(saved_path)
                add_record("documents", {"title": title.strip(), "category": category, "context": context.strip(), "path": path_value.strip(), "notes": notes.strip()})
                rerun_notice("Documento archiviato")


def render(module):
    module_header(module["name"], module.get("description", ""), "DOCUMENTI OPERATIVI")
    templates = _template_files()
    archive_tab, templates_tab, compiler_tab, upload_tab = st.tabs([
        "Archivio", "Modelli VA Digital", "Compila modello", "Carica documento"
    ])

    with archive_tab:
        _render_archive()
    with templates_tab:
        _render_templates(templates)
    with compiler_tab:
        _render_compiler(templates)
    with upload_tab:
        _render_upload()
