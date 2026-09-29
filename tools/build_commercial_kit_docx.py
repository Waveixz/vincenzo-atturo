from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "output" / "commercial-kit"

BLUE = "168FF0"
NAVY = "0B2C47"
PALE = "EAF4FC"
LIGHT = "F4F7FA"
MID = "60748A"
BORDER = "D9E2EA"
BLACK = "111820"
WHITE = "FFFFFF"


def set_cell_shading(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_cell_border(cell, color=BORDER, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=110, start=130, bottom=110, end=130):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def setup_document(title, subtitle):
    doc = Document()
    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(18)
    section.bottom_margin = Mm(17)
    section.left_margin = Mm(19)
    section.right_margin = Mm(19)
    section.header_distance = Mm(7)
    section.footer_distance = Mm(7)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.2)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12

    title_style = styles["Title"]
    title_style.font.name = "Bahnschrift"
    title_style.font.size = Pt(24)
    title_style.font.bold = False
    title_style.font.color.rgb = RGBColor.from_string("000000")
    title_style.paragraph_format.space_after = Pt(5)

    for style_name, size, before, after in (
        ("Heading 1", 15, 14, 6),
        ("Heading 2", 11.5, 10, 4),
    ):
        style = styles[style_name]
        style.font.name = "Bahnschrift"
        style.font.size = Pt(size)
        style.font.bold = False
        style.font.color.rgb = RGBColor.from_string("000000")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header
    table = header.add_table(rows=1, cols=2, width=Mm(172))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.columns[0].width = Mm(74)
    table.columns[1].width = Mm(98)
    left, right = table.rows[0].cells
    left_p = left.paragraphs[0]
    left_p.paragraph_format.space_after = Pt(0)
    run = left_p.add_run("VA")
    run.font.name = "Bahnschrift"
    run.font.size = Pt(17)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(BLUE)
    run = left_p.add_run("  VA DIGITAL")
    run.font.name = "Calibri"
    run.font.size = Pt(8.5)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(NAVY)

    right_p = right.paragraphs[0]
    right_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    right_p.paragraph_format.space_after = Pt(0)
    run = right_p.add_run("va-digital.it  |  +39 334 991 6514")
    run.font.name = "Calibri"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(MID)
    for cell in (left, right):
        set_cell_border(cell, BLUE, "8")
        set_cell_margins(cell, top=0, bottom=80, start=180, end=180)

    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    run = p.add_run("VA Digital  |  Vincenzo Atturo  |  atturo.vincenzo@gmail.com")
    run.font.name = "Calibri"
    run.font.size = Pt(7.5)
    run.font.color.rgb = RGBColor.from_string(MID)

    p = doc.add_paragraph(style="Title")
    p.add_run(title)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(12)
    run = p.add_run(subtitle)
    run.font.size = Pt(10.5)
    run.font.color.rgb = RGBColor.from_string(MID)
    return doc


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    keep_with_next(p)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Mm(5)
        p.paragraph_format.first_line_indent = Mm(-2.5)
        p.paragraph_format.space_after = Pt(3)
        p.add_run(item)


def add_label_line(doc, label, hint=""):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(label + "  ")
    r.bold = True
    if hint:
        r = p.add_run(hint)
        r.italic = True
        r.font.color.rgb = RGBColor.from_string(MID)
    p.add_run("\n" + "_" * 92)
    return p


def add_table(doc, headers, rows, widths=None, font_size=9):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    if widths:
        for i, width in enumerate(widths):
            table.columns[i].width = Mm(width)
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, text in enumerate(headers):
        cell = hdr.cells[i]
        set_cell_shading(cell, NAVY)
        set_cell_border(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(text)
        run.bold = True
        run.font.color.rgb = RGBColor.from_string(WHITE)
        run.font.size = Pt(font_size)
    for row_index, row_data in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row_data):
            cell = cells[i]
            set_cell_shading(cell, WHITE if row_index % 2 == 0 else LIGHT)
            set_cell_border(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            run = p.add_run(str(value))
            run.font.size = Pt(font_size)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def save(doc, filename):
    path = OUT / filename
    doc.save(path)
    return path


def build_company_sheet():
    doc = setup_document(
        "Presentazione VA Digital",
        "Soluzioni software, automazioni, dashboard e GIS progettate intorno ai processi reali.",
    )
    add_heading(doc, "Profilo", 1)
    doc.add_paragraph(
        "VA Digital affianca aziende, professionisti ed enti nella progettazione di strumenti digitali su misura. "
        "L'attività parte dal processo esistente, individua passaggi ripetitivi o dati dispersi e costruisce una soluzione utilizzabile nel lavoro quotidiano."
    )
    add_heading(doc, "Ambiti di intervento", 1)
    add_table(
        doc,
        ["Ambito", "Applicazioni"],
        [
            ["Software su misura", "Applicazioni web, strumenti operativi e portali dedicati"],
            ["Automazioni", "Workflow, script Python, report e integrazioni tra sistemi"],
            ["Dati e dashboard", "KPI, database, processi ETL e viste operative"],
            ["GIS e territorio", "WebGIS, cartografia, dati catastali e territoriali"],
        ],
        widths=[45, 127],
        font_size=9.2,
    )
    add_heading(doc, "Metodo di lavoro", 1)
    add_bullets(
        doc,
        [
            "Analisi del processo e delle persone che lo utilizzano.",
            "Definizione di un primo perimetro verificabile.",
            "Sviluppo per fasi con verifiche intermedie e consegna documentata.",
        ],
    )
    add_heading(doc, "Contatti", 1)
    p = doc.add_paragraph()
    p.add_run("Vincenzo Atturo").bold = True
    p.add_run("\nIT Manager e Digital Solutions Developer")
    p.add_run("\n+39 334 991 6514  |  atturo.vincenzo@gmail.com  |  va-digital.it")
    return save(doc, "01_scheda_presentazione_va_digital.docx")


def build_requirements_form():
    doc = setup_document(
        "Raccolta esigenze e requisiti",
        "Modulo iniziale per comprendere il processo, il problema da risolvere e il risultato atteso.",
    )
    add_heading(doc, "Dati del cliente", 1)
    add_table(
        doc,
        ["Informazione", "Da compilare"],
        [
            ["Azienda o ente", ""],
            ["Referente e ruolo", ""],
            ["Email e telefono", ""],
            ["Data del confronto", ""],
        ],
        widths=[52, 120],
    )
    add_heading(doc, "Contesto", 1)
    add_label_line(doc, "Processo o attività interessata", "Descrivere brevemente cosa accade oggi")
    add_label_line(doc, "Problema principale", "Tempi, errori, passaggi manuali, dati dispersi o altro")
    add_label_line(doc, "Risultato atteso", "Indicare come dovrebbe funzionare il processo dopo l'intervento")
    add_heading(doc, "Utenti e utilizzo", 1)
    add_table(
        doc,
        ["Domanda", "Risposta"],
        [
            ["Chi utilizzerà la soluzione", ""],
            ["Numero indicativo di utenti", ""],
            ["Frequenza di utilizzo", "[ ] Giornaliera   [ ] Settimanale   [ ] Occasionale"],
            ["Accesso", "[ ] Interno   [ ] Da remoto   [ ] Pubblico"],
        ],
        widths=[60, 112],
    )
    add_heading(doc, "Dati e integrazioni", 1)
    add_bullets(
        doc,
        [
            "Fonti disponibili: file Excel, database, email, API, documenti, dati GIS o altro.",
            "Sistemi con cui la soluzione deve comunicare.",
            "Dati personali, riservati o soggetti a regole specifiche.",
        ],
    )
    add_label_line(doc, "Note tecniche e vincoli")
    add_heading(doc, "Priorità", 1)
    add_table(
        doc,
        ["Elemento", "Valutazione"],
        [
            ["Urgenza", "[ ] Alta   [ ] Media   [ ] Bassa"],
            ["Budget indicativo", "€ ____________________"],
            ["Data desiderata", "____________________"],
            ["Referente per le decisioni", "____________________"],
        ],
        widths=[60, 112],
    )
    add_heading(doc, "Criteri di successo", 1)
    add_label_line(doc, "Come misurare il risultato", "Tempo risparmiato, errori ridotti, dati disponibili o altro")
    add_label_line(doc, "Condizione minima di accettazione", "Indicare cosa deve funzionare alla prima consegna")
    add_heading(doc, "Vincoli e approvazioni", 1)
    add_table(
        doc,
        ["Domanda", "Risposta"],
        [
            ["Accessi o autorizzazioni necessari", ""],
            ["Regole di sicurezza o riservatezza", ""],
            ["Chi verifica il risultato", ""],
            ["Chi approva l'avvio", ""],
        ],
        widths=[72, 100],
    )
    add_heading(doc, "Note del confronto", 1)
    add_label_line(doc, "Decisioni, dubbi e prossimi passi")
    return save(doc, "02_modulo_requisiti_cliente.docx")


def build_proposal():
    doc = setup_document(
        "Preventivo e proposta tecnica",
        "Documento per definire obiettivi, attività, tempi, costi e condizioni del progetto.",
    )
    add_table(
        doc,
        ["Riferimento", "Valore"],
        [
            ["Cliente", "[Ragione sociale o nominativo]"],
            ["Progetto", "[Titolo del progetto]"],
            ["Numero proposta", "[VA-AAAA-NNN]"],
            ["Data e validità", "[Data] - valida fino al [Data]"],
        ],
        widths=[50, 122],
    )
    add_heading(doc, "Esigenza", 1)
    doc.add_paragraph("[Descrivere il contesto del cliente, il problema rilevato e l'obiettivo concordato.]")
    add_heading(doc, "Soluzione proposta", 1)
    doc.add_paragraph("[Descrivere la soluzione, gli utenti interessati e il risultato operativo atteso.]")
    add_heading(doc, "Perimetro delle attività", 1)
    add_table(
        doc,
        ["Fase", "Attività", "Risultato"],
        [
            ["1 Analisi", "[Attività incluse]", "[Documento o decisione]"],
            ["2 Realizzazione", "[Attività incluse]", "[Modulo o funzionalità]"],
            ["3 Verifica", "[Test e correzioni]", "[Versione approvata]"],
            ["4 Consegna", "[Pubblicazione e passaggio]", "[Materiali consegnati]"],
        ],
        widths=[30, 82, 60],
    )
    add_heading(doc, "Tempi", 1)
    doc.add_paragraph("Avvio previsto: [Data]   Durata stimata: [Numero settimane]   Consegna prevista: [Data]")
    add_heading(doc, "Corrispettivo", 1)
    add_table(
        doc,
        ["Voce", "Importo"],
        [
            ["Analisi e progettazione", "€ [Importo]"],
            ["Sviluppo e configurazione", "€ [Importo]"],
            ["Test, consegna e affiancamento", "€ [Importo]"],
            ["Totale", "€ [Importo] oltre eventuali oneri applicabili"],
        ],
        widths=[120, 52],
    )
    add_heading(doc, "Condizioni", 1)
    add_bullets(
        doc,
        [
            "Pagamento: [Modalità e scadenze].",
            "Attività escluse: [Elenco delle esclusioni].",
            "Le richieste fuori perimetro richiedono una valutazione separata.",
            "Il cliente mette a disposizione accessi, dati e referenti necessari nei tempi concordati.",
        ],
    )
    add_heading(doc, "Accettazione", 1)
    add_table(
        doc,
        ["Per VA Digital", "Per il cliente"],
        [["Vincenzo Atturo\nData e firma", "[Nome e ruolo]\nData e firma"]],
        widths=[86, 86],
    )
    return save(doc, "03_preventivo_proposta_tecnica.docx")


def build_kickoff():
    doc = setup_document(
        "Documento di avvio progetto",
        "Riferimento condiviso per obiettivi, perimetro, tempi, responsabilità e modalità operative.",
    )
    add_table(
        doc,
        ["Dato", "Valore"],
        [
            ["Cliente", "[Cliente]"],
            ["Progetto", "[Titolo]"],
            ["Data di avvio", "[Data]"],
            ["Referenti", "VA Digital: Vincenzo Atturo\nCliente: [Nome e ruolo]"],
        ],
        widths=[50, 122],
    )
    add_heading(doc, "Obiettivo", 1)
    doc.add_paragraph("[Descrivere il risultato che il progetto deve rendere disponibile e il problema operativo che affronta.]")
    add_heading(doc, "Perimetro iniziale", 1)
    add_bullets(doc, ["Incluso: [Attività e funzionalità]", "Escluso: [Attività non comprese]", "Dipendenze: [Accessi, dati e decisioni necessarie]"])
    add_heading(doc, "Piano di lavoro", 1)
    add_table(
        doc,
        ["Fase", "Periodo", "Responsabile", "Criterio di completamento"],
        [
            ["Analisi", "[Date]", "[Nome]", "Requisiti confermati"],
            ["Realizzazione", "[Date]", "[Nome]", "Funzioni concordate disponibili"],
            ["Verifica", "[Date]", "[Nome]", "Test completati"],
            ["Consegna", "[Data]", "[Nome]", "Verbale firmato"],
        ],
        widths=[32, 38, 42, 60],
    )
    add_heading(doc, "Responsabilità", 1)
    add_table(
        doc,
        ["Soggetto", "Responsabilità"],
        [
            ["VA Digital", "Progettazione, realizzazione, verifiche tecniche e documentazione concordata"],
            ["Cliente", "Accessi, dati, decisioni, verifiche funzionali e approvazioni"],
        ],
        widths=[42, 130],
    )
    add_heading(doc, "Comunicazione e decisioni", 1)
    doc.add_paragraph("Canale principale: [Email o strumento]   Frequenza aggiornamenti: [Frequenza]   Tempo di risposta atteso: [Tempo]")
    add_label_line(doc, "Decisioni prese durante l'avvio")
    return save(doc, "04_avvio_progetto.docx")


def build_status_report():
    doc = setup_document(
        "Report sullo stato dei lavori",
        "Aggiornamento periodico su attività completate, prossime scadenze, decisioni e criticità.",
    )
    add_table(
        doc,
        ["Dato", "Valore"],
        [
            ["Cliente e progetto", "[Cliente] - [Progetto]"],
            ["Periodo", "[Dal] - [Al]"],
            ["Responsabile", "[Nome]"],
            ["Stato complessivo", "[ ] Regolare   [ ] Attenzione   [ ] Critico"],
        ],
        widths=[50, 122],
    )
    add_heading(doc, "Sintesi", 1)
    doc.add_paragraph("[Indicare cosa è cambiato nel periodo, cosa è stato completato e quale decisione o supporto serve ora.]")
    add_heading(doc, "Avanzamento", 1)
    add_table(
        doc,
        ["Attività", "Stato", "Evidenza o risultato", "Scadenza"],
        [
            ["[Attività 1]", "[Completata/In corso]", "[Risultato]", "[Data]"],
            ["[Attività 2]", "[Completata/In corso]", "[Risultato]", "[Data]"],
            ["[Attività 3]", "[Da avviare]", "[Risultato atteso]", "[Data]"],
        ],
        widths=[48, 30, 66, 28],
        font_size=8.6,
    )
    add_heading(doc, "Prossime attività", 1)
    add_bullets(doc, ["[Attività e responsabile]", "[Attività e responsabile]", "[Attività e responsabile]"])
    add_heading(doc, "Decisioni richieste", 1)
    add_table(
        doc,
        ["Decisione", "Responsabile", "Entro"],
        [["[Decisione necessaria]", "[Nome]", "[Data]"], ["[Decisione necessaria]", "[Nome]", "[Data]"]],
        widths=[100, 44, 28],
    )
    add_heading(doc, "Rischi e impedimenti", 1)
    add_table(
        doc,
        ["Elemento", "Impatto", "Azione"],
        [["[Rischio o blocco]", "[Basso/Medio/Alto]", "[Azione e responsabile]"]],
        widths=[70, 36, 66],
    )
    return save(doc, "05_report_stato_lavori.docx")


def build_delivery():
    doc = setup_document(
        "Consegna e collaudo",
        "Verbale per registrare materiali consegnati, verifiche eseguite, esiti e accettazione del cliente.",
    )
    doc.sections[0].bottom_margin = Mm(11)
    add_table(
        doc,
        ["Dato", "Valore"],
        [
            ["Cliente", "[Cliente]"],
            ["Progetto", "[Progetto]"],
            ["Versione", "[Versione o data]"],
            ["Data di consegna", "[Data]"],
        ],
        widths=[50, 122],
    )
    add_heading(doc, "Materiali consegnati", 1)
    add_table(
        doc,
        ["Elemento", "Versione o percorso", "Esito"],
        [
            ["Applicazione o configurazione", "[Riferimento]", "[ ] Consegnato"],
            ["Documentazione", "[Riferimento]", "[ ] Consegnato"],
            ["Credenziali o accessi", "[Canale separato]", "[ ] Consegnato"],
            ["Backup o esportazione", "[Riferimento]", "[ ] Consegnato"],
        ],
        widths=[64, 70, 38],
    )
    add_heading(doc, "Verifiche di collaudo", 1)
    add_table(
        doc,
        ["Verifica", "Risultato", "Note"],
        [
            ["Accesso e permessi", "[ ] Conforme   [ ] Da correggere", ""],
            ["Funzioni concordate", "[ ] Conforme   [ ] Da correggere", ""],
            ["Dati e integrazioni", "[ ] Conforme   [ ] Da correggere", ""],
            ["Prestazioni e compatibilità", "[ ] Conforme   [ ] Da correggere", ""],
        ],
        widths=[58, 66, 48],
        font_size=8.7,
    )
    add_heading(doc, "Osservazioni e attività residue", 1)
    add_label_line(doc, "Elementi da completare o correggere")
    add_heading(doc, "Esito", 1)
    doc.add_paragraph("[ ] Accettato senza riserve   [ ] Accettato con attività residue   [ ] Collaudo da ripetere")
    signature_table = add_table(
        doc,
        ["Per VA Digital", "Per il cliente"],
        [["Vincenzo Atturo\nData e firma", "[Nome e ruolo]\nData e firma"]],
        widths=[86, 86],
    )
    # Word richiede sempre un paragrafo dopo una tabella: lo rendiamo minimo
    # per evitare che venga spinto da solo su una seconda pagina vuota.
    trailing = signature_table._tbl.getnext()
    if trailing is not None:
        trailing_p = doc.paragraphs[-1]
        trailing_p.paragraph_format.space_before = Pt(0)
        trailing_p.paragraph_format.space_after = Pt(0)
        trailing_p.paragraph_format.line_spacing = Pt(1)
        run = trailing_p.add_run("")
        run.font.size = Pt(1)
    return save(doc, "06_consegna_collaudo.docx")


def build_support_sheet():
    doc = setup_document(
        "Assistenza e manutenzione",
        "Scheda per definire copertura, priorità, canali di richiesta e attività comprese dopo la consegna.",
    )
    add_table(
        doc,
        ["Dato", "Valore"],
        [
            ["Cliente e soluzione", "[Cliente] - [Soluzione]"],
            ["Periodo di copertura", "[Dal] - [Al]"],
            ["Referente", "[Nome e contatto]"],
            ["Canale richieste", "[Email o sistema concordato]"],
        ],
        widths=[54, 118],
    )
    add_heading(doc, "Copertura", 1)
    add_bullets(
        doc,
        [
            "Correzione di anomalie riproducibili nelle funzioni consegnate.",
            "Aggiornamenti tecnici concordati e verifiche periodiche.",
            "Supporto operativo entro i limiti del pacchetto scelto.",
        ],
    )
    add_heading(doc, "Livelli di priorità", 1)
    add_table(
        doc,
        ["Priorità", "Descrizione", "Presa in carico prevista"],
        [
            ["Critica", "Servizio non utilizzabile o perdita di una funzione essenziale", "[Tempo]"],
            ["Alta", "Funzione importante compromessa con alternativa limitata", "[Tempo]"],
            ["Ordinaria", "Anomalia non bloccante o richiesta di chiarimento", "[Tempo]"],
            ["Evolutiva", "Nuova funzione o modifica del perimetro", "Valutazione separata"],
        ],
        widths=[28, 100, 44],
        font_size=8.5,
    )
    add_heading(doc, "Piano scelto", 1)
    add_table(
        doc,
        ["Voce", "Condizione"],
        [
            ["Canone o pacchetto ore", "[Importo o ore]"],
            ["Ore incluse", "[Numero]"],
            ["Frequenza controllo", "[Mensile/Trimestrale/Altro]"],
            ["Attività escluse", "[Elenco]"],
        ],
        widths=[60, 112],
    )
    add_heading(doc, "Come aprire una richiesta", 1)
    doc.add_paragraph("Indicare cliente, sistema interessato, descrizione del problema, passaggi per riprodurlo, urgenza e schermate utili. Non inviare password nello stesso messaggio.")
    return save(doc, "07_assistenza_manutenzione.docx")


def build_letterhead():
    doc = setup_document(
        "Carta intestata VA Digital",
        "Modello per comunicazioni, dichiarazioni e documenti brevi destinati a clienti e fornitori.",
    )
    add_table(
        doc,
        ["Riferimento", "Valore"],
        [
            ["Destinatario", "[Nome, azienda o ente]"],
            ["Oggetto", "[Oggetto della comunicazione]"],
            ["Data", "[Data]"],
        ],
        widths=[44, 128],
    )
    doc.add_paragraph()
    doc.add_paragraph("Gentile [Nome],")
    doc.add_paragraph("[Inserire qui il testo della comunicazione. Utilizzare paragrafi brevi e indicare con chiarezza richieste, decisioni o prossimi passaggi.]")
    doc.add_paragraph("[Secondo paragrafo facoltativo.]")
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.add_run("Cordiali saluti\n").bold = True
    p.add_run("Vincenzo Atturo\nIT Manager e Digital Solutions Developer\nVA Digital")
    return save(doc, "08_carta_intestata_va_digital.docx")


def build_service_contract():
    doc = setup_document(
        "Contratto di incarico per servizi digitali",
        "Modello B2B da adattare al singolo progetto e da sottoporre a verifica professionale prima della firma.",
    )
    doc.styles["Normal"].font.size = Pt(9.4)
    doc.styles["Normal"].paragraph_format.space_after = Pt(4.5)
    doc.styles["Normal"].paragraph_format.line_spacing = 1.08

    add_heading(doc, "Dati del contratto", 1)
    add_table(
        doc,
        ["Riferimento", "Informazione da completare"],
        [
            ["Numero e data", "[VA-AAAA-NNN] - [Data]"],
            ["Progetto", "[Titolo del progetto]"],
            ["Preventivo collegato", "[Numero e data del preventivo]"],
            ["Durata prevista", "[Data di avvio] - [Data o durata stimata]"],
        ],
        widths=[48, 124],
        font_size=8.8,
    )

    add_heading(doc, "Parti", 1)
    doc.add_paragraph(
        "Tra [denominazione o nome del fornitore], con sede o domicilio in [indirizzo], codice fiscale e partita IVA "
        "[dati], email [email], PEC [PEC], rappresentato da [nome e ruolo], di seguito Fornitore,"
    )
    doc.add_paragraph("e")
    doc.add_paragraph(
        "[denominazione del cliente], con sede in [indirizzo], codice fiscale e partita IVA [dati], email [email], "
        "PEC [PEC], rappresentato da [nome e ruolo], di seguito Cliente."
    )
    doc.add_paragraph("Fornitore e Cliente sono congiuntamente indicati come Parti.")

    add_heading(doc, "Premesse e documenti contrattuali", 1)
    doc.add_paragraph(
        "Il Cliente intende affidare al Fornitore le attività descritte nel presente contratto. Le premesse e gli allegati "
        "ne costituiscono parte integrante. In caso di contrasto prevalgono, nell'ordine: il presente contratto, il preventivo "
        "accettato, la specifica tecnica o il documento dei requisiti, il documento di avvio progetto e gli altri allegati."
    )

    add_heading(doc, "1 Oggetto dell'incarico", 1)
    doc.add_paragraph(
        "Il Fornitore realizza i servizi indicati nel preventivo e negli allegati, che possono comprendere analisi, sviluppo "
        "software, automazioni, dashboard, integrazioni, elaborazioni dati, soluzioni GIS, configurazione, documentazione, "
        "formazione e assistenza. Sono comprese esclusivamente le attività espressamente elencate."
    )

    add_heading(doc, "2 Perimetro e variazioni", 1)
    doc.add_paragraph(
        "Requisiti, risultati attesi, esclusioni e criteri di accettazione sono riportati negli allegati. Ogni richiesta che "
        "modifichi funzionalità, dati, integrazioni, utenti, ambienti o tempi è valutata separatamente. Il Fornitore comunica "
        "l'eventuale impatto su costo e calendario e procede dopo approvazione scritta del Cliente."
    )

    add_heading(doc, "3 Modalità di esecuzione e collaborazione", 1)
    doc.add_paragraph(
        "Il Fornitore opera con autonomia organizzativa e senza vincolo di subordinazione. Il Cliente mette a disposizione "
        "referenti, accessi, dati, ambienti, decisioni e riscontri necessari. Le stime presuppongono che tali elementi siano "
        "forniti nei tempi concordati. Ritardi o indisponibilità imputabili al Cliente possono comportare l'aggiornamento del calendario."
    )

    add_heading(doc, "4 Tempi e avanzamento", 1)
    doc.add_paragraph(
        "L'avvio è previsto per [data], con durata stimata di [periodo]. Le scadenze intermedie sono indicate nel piano di "
        "lavoro. Salvo espressa qualificazione scritta come essenziale, le date hanno natura previsionale e possono essere "
        "aggiornate per variazioni approvate, dipendenze esterne o ritardi nella collaborazione del Cliente."
    )

    add_heading(doc, "5 Corrispettivo, fatturazione e pagamenti", 1)
    add_table(
        doc,
        ["Voce", "Condizione da definire"],
        [
            ["Corrispettivo", "Euro [importo] oltre imposte e oneri applicabili"],
            ["Piano di pagamento", "[Acconto] - [stato avanzamento] - [saldo]"],
            ["Termine di pagamento", "[Numero] giorni dalla data della fattura"],
            ["Spese", "[Incluse / rimborsate previa autorizzazione]"],
        ],
        widths=[48, 124],
        font_size=8.8,
    )
    doc.add_paragraph(
        "In caso di ritardo si applicano gli interessi previsti dalla normativa vigente. Dopo [numero] giorni dal sollecito, "
        "il Fornitore può sospendere le attività, previa comunicazione scritta, fermo il pagamento delle prestazioni già eseguite."
    )

    add_heading(doc, "6 Verifica, collaudo e accettazione", 1)
    doc.add_paragraph(
        "Alla consegna il Cliente dispone di [10] giorni lavorativi per eseguire le verifiche rispetto ai criteri concordati. "
        "Eventuali difformità devono essere comunicate per iscritto con descrizione riproducibile. Il Fornitore corregge le "
        "difformità comprese nel perimetro. In assenza di rilievi sostanziali entro il termine, oppure con l'uso in produzione, "
        "la consegna si considera accettata, salvo vizi non riconoscibili con l'ordinaria verifica."
    )

    add_heading(doc, "7 Proprietà intellettuale e licenze", 1)
    doc.add_paragraph(
        "Ciascuna Parte conserva la titolarità dei materiali, dati, strumenti, modelli e componenti preesistenti. Componenti "
        "generici, librerie, procedure, metodi e know-how riutilizzabili del Fornitore restano di sua titolarità. Software e "
        "contenuti di terzi restano soggetti alle rispettive licenze."
    )
    doc.add_paragraph(
        "Per i risultati sviluppati specificamente per il Cliente, scegliere una sola opzione e completarla: "
        "[A: cessione dei diritti patrimoniali dopo il pagamento integrale] oppure [B: licenza d'uso non esclusiva, "
        "illimitata nel tempo e destinata alle attività del Cliente] oppure [C: condizioni indicate nell'allegato tecnico]. "
        "La consegna del codice sorgente è inclusa solo se espressamente prevista."
    )

    add_heading(doc, "8 Dati, contenuti e copie di sicurezza", 1)
    doc.add_paragraph(
        "Il Cliente garantisce di poter utilizzare e comunicare i dati e i contenuti forniti e ne verifica correttezza, qualità "
        "e completezza. Salvo diverso accordo, il Cliente conserva una copia aggiornata dei dati originali e dei materiali "
        "caricati. Migrazione, bonifica, conservazione e backup continuativo sono inclusi solo se indicati nel perimetro."
    )

    add_heading(doc, "9 Riservatezza", 1)
    doc.add_paragraph(
        "Le Parti mantengono riservate le informazioni tecniche, commerciali e organizzative ricevute in occasione dell'incarico "
        "e le utilizzano soltanto per eseguirlo. L'obbligo non riguarda informazioni già pubbliche, legittimamente conosciute o "
        "la cui comunicazione sia richiesta dalla legge. L'obbligo prosegue per [3] anni dalla cessazione del contratto."
    )

    add_heading(doc, "10 Protezione dei dati personali", 1)
    doc.add_paragraph(
        "Le Parti trattano i dati dei rispettivi referenti come autonomi titolari per finalità connesse al contratto. Se il "
        "Fornitore tratta dati personali per conto del Cliente, le Parti sottoscrivono prima dell'avvio un accordo conforme "
        "all'articolo 28 del Regolamento UE 2016/679, con istruzioni, durata, finalità, categorie di dati, misure di sicurezza "
        "e condizioni per eventuali sub-responsabili. Tale accordo prevale per gli aspetti relativi al trattamento affidato."
    )

    add_heading(doc, "11 Accessi e sicurezza", 1)
    doc.add_paragraph(
        "Gli accessi sono nominativi ove possibile e limitati alle attività necessarie. Le credenziali sono trasmesse con canale "
        "separato e non sono inserite nel contratto. Il Cliente revoca o modifica gli accessi alla conclusione dell'incarico. "
        "Misure, ambienti e responsabilità operative ulteriori sono indicati nell'allegato tecnico."
    )

    add_heading(doc, "12 Garanzia, assistenza e manutenzione", 1)
    doc.add_paragraph(
        "Per [60] giorni dall'accettazione il Fornitore corregge senza costi aggiuntivi le anomalie riproducibili che impediscono "
        "alle funzioni consegnate di rispettare i requisiti approvati. Sono escluse nuove esigenze, modifiche di terzi, uso non "
        "conforme, interventi non autorizzati e problemi di infrastrutture non gestite dal Fornitore. Assistenza evolutiva e "
        "manutenzione successiva richiedono un accordo separato."
    )

    add_heading(doc, "13 Responsabilità", 1)
    doc.add_paragraph(
        "Ciascuna Parte risponde dei danni diretti derivanti da propri inadempimenti secondo la legge applicabile. Salvo dolo, "
        "colpa grave e limiti inderogabili di legge, la responsabilità complessiva del Fornitore è limitata a [importo oppure "
        "corrispettivo pagato per il progetto]. Restano esclusi, nei limiti consentiti, danni indiretti, perdita di profitto, "
        "interruzione dell'attività o perdita di dati evitabile mediante copie di sicurezza previste."
    )

    add_heading(doc, "14 Sospensione, recesso e risoluzione", 1)
    doc.add_paragraph(
        "Ciascuna Parte può risolvere il contratto per grave inadempimento non sanato entro [15] giorni dal ricevimento di una "
        "contestazione scritta. Il Cliente può recedere con preavviso di [15] giorni, corrispondendo le attività eseguite, gli "
        "impegni già assunti e gli eventuali costi di chiusura concordati. Il Fornitore può sospendere le attività in caso di "
        "mancato pagamento, accessi non disponibili o istruzioni che espongano a rischi illegittimi, dopo avviso scritto."
    )

    add_heading(doc, "15 Forza maggiore", 1)
    doc.add_paragraph(
        "Nessuna Parte risponde di ritardi causati da eventi imprevedibili e fuori dal proprio ragionevole controllo. La Parte "
        "interessata informa tempestivamente l'altra e adotta misure ragionevoli per limitarne gli effetti. Se l'impedimento "
        "prosegue oltre [30] giorni, le Parti valutano la modifica o la cessazione dell'incarico."
    )

    add_heading(doc, "16 Collaboratori e servizi di terzi", 1)
    doc.add_paragraph(
        "Il Fornitore può avvalersi di collaboratori qualificati, restando responsabile delle attività affidate. Servizi cloud, "
        "API, software o licenze di terzi sono soggetti a condizioni, disponibilità e costi dei rispettivi fornitori. Eventuali "
        "sub-responsabili del trattamento sono gestiti secondo l'accordo privacy applicabile."
    )

    add_heading(doc, "17 Comunicazioni", 1)
    doc.add_paragraph(
        "Comunicazioni operative e approvazioni sono inviate agli indirizzi indicati nel documento di avvio progetto. Diffide, "
        "recesso e risoluzione sono inviati tramite PEC o altro mezzo idoneo a provarne ricezione e contenuto."
    )

    add_heading(doc, "18 Legge applicabile e foro", 1)
    doc.add_paragraph(
        "Il contratto è regolato dalla legge italiana. Le Parti cercano prima una soluzione bonaria entro [30] giorni dalla "
        "contestazione. Per le controversie non risolte è competente in via esclusiva il Foro di [città], salvo competenze "
        "inderogabili. Questa clausola è predisposta per rapporti B2B e deve essere adattata se il Cliente è un consumatore o un ente pubblico."
    )

    add_heading(doc, "19 Disposizioni finali", 1)
    doc.add_paragraph(
        "Modifiche e integrazioni sono valide se approvate per iscritto. L'eventuale invalidità di una clausola non pregiudica "
        "le altre. Il mancato esercizio di un diritto non costituisce rinuncia. Il contratto sostituisce le intese precedenti "
        "sullo stesso oggetto, salvo gli allegati espressamente richiamati."
    )

    add_heading(doc, "Allegati", 1)
    add_bullets(
        doc,
        [
            "Allegato A - Preventivo e proposta tecnica.",
            "Allegato B - Requisiti, perimetro e criteri di accettazione.",
            "Allegato C - Piano di lavoro e responsabilità.",
            "Allegato D - Accordo sul trattamento dei dati personali, se applicabile.",
            "Allegato E - Assistenza e manutenzione, se prevista.",
        ],
    )

    add_heading(doc, "Sottoscrizione", 1)
    add_table(
        doc,
        ["Per il Fornitore", "Per il Cliente"],
        [["Nome: [nome]\nData: [data]\nFirma: ____________________", "Nome e ruolo: [dati]\nData: [data]\nFirma: ____________________"]],
        widths=[86, 86],
        font_size=8.8,
    )

    add_heading(doc, "Approvazione specifica", 1)
    doc.add_paragraph(
        "Ai sensi e per gli effetti degli articoli 1341 e 1342 del codice civile, il Cliente dichiara di approvare "
        "specificamente, dopo attenta lettura, le clausole: 4 Tempi e avanzamento; 5 Corrispettivo, fatturazione e pagamenti; "
        "6 Verifica, collaudo e accettazione; 7 Proprietà intellettuale e licenze; 8 Dati, contenuti e copie di sicurezza; "
        "12 Garanzia, assistenza e manutenzione; 13 Responsabilità; 14 Sospensione, recesso e risoluzione; 16 Collaboratori e "
        "servizi di terzi; 18 Legge applicabile e foro."
    )
    add_table(
        doc,
        ["Per il Cliente"],
        [["Nome e ruolo: [dati]     Data: [data]     Firma: ______________________________"]],
        widths=[172],
        font_size=8.8,
    )
    doc.add_paragraph(
        "Nota d'uso: questo modello è uno strumento operativo generale e non sostituisce la consulenza legale, fiscale o "
        "privacy. Prima dell'impiego eliminare le alternative non scelte, completare tutti i campi e verificare coerenza con "
        "preventivo, regime professionale, coperture assicurative, tipo di cliente e trattamento effettivo dei dati."
    )
    return save(doc, "10_contratto_incarico_servizi_digitali.docx")


def build_data_processing_agreement():
    doc = setup_document(
        "Accordo sul trattamento dei dati personali",
        "Modello di nomina e istruzioni al responsabile del trattamento ai sensi dell'articolo 28 del Regolamento UE 2016 679.",
    )
    doc.styles["Normal"].font.size = Pt(9.2)
    doc.styles["Normal"].paragraph_format.space_after = Pt(4.2)
    doc.styles["Normal"].paragraph_format.line_spacing = 1.06

    add_heading(doc, "Dati dell'accordo", 1)
    add_table(
        doc,
        ["Riferimento", "Informazione da completare"],
        [
            ["Contratto principale", "[Numero, data e progetto]"],
            ["Titolare del trattamento", "[Cliente, sede, dati fiscali e referente privacy]"],
            ["Responsabile del trattamento", "[Fornitore, sede, dati fiscali e referente privacy]"],
            ["Durata", "[Durata del servizio e periodo di restituzione o cancellazione]"],
        ],
        widths=[52, 120],
        font_size=8.6,
    )

    add_heading(doc, "Premesse e ruolo delle parti", 1)
    doc.add_paragraph(
        "Il Titolare determina finalità e mezzi essenziali dei trattamenti descritti nel presente accordo. Il Responsabile "
        "tratta i dati personali esclusivamente per conto del Titolare, nei limiti del contratto principale e delle istruzioni "
        "documentate. Il presente accordo integra il contratto principale e prevale per gli aspetti relativi ai trattamenti affidati."
    )

    add_heading(doc, "1 Oggetto, natura e finalità", 1)
    add_table(
        doc,
        ["Elemento", "Descrizione da completare"],
        [
            ["Servizi interessati", "[Sviluppo, hosting, assistenza, migrazione, analisi o altro]"],
            ["Operazioni", "[Raccolta, consultazione, organizzazione, modifica, esportazione, cancellazione o altro]"],
            ["Finalità", "[Finalità operative definite dal Titolare]"],
            ["Luoghi e sistemi", "[Infrastrutture, Paesi e servizi cloud utilizzati]"],
        ],
        widths=[48, 124],
        font_size=8.5,
    )

    add_heading(doc, "2 Categorie di interessati e dati", 1)
    add_table(
        doc,
        ["Voce", "Selezionare e specificare"],
        [
            ["Interessati", "[Clienti] [dipendenti] [fornitori] [utenti] [cittadini] [altro]"],
            ["Dati comuni", "[Identificativi] [contatti] [account] [dati professionali] [log] [altro]"],
            ["Categorie particolari", "[Non previste / descrivere dati ex articolo 9]"],
            ["Dati giudiziari", "[Non previsti / descrivere dati ex articolo 10]"],
            ["Volume e frequenza", "[Occasionale / continuativo, quantità indicativa]"],
        ],
        widths=[48, 124],
        font_size=8.4,
    )

    add_heading(doc, "3 Istruzioni documentate", 1)
    doc.add_paragraph(
        "Il Responsabile tratta i dati soltanto su istruzione documentata del Titolare, anche per eventuali trasferimenti "
        "verso Paesi terzi, salvo obblighi di legge applicabili. Se ritiene che un'istruzione violi la normativa, informa "
        "tempestivamente il Titolare e sospende l'attività interessata in attesa di chiarimenti, quando consentito."
    )

    add_heading(doc, "4 Persone autorizzate e riservatezza", 1)
    doc.add_paragraph(
        "Il Responsabile consente l'accesso ai dati soltanto a persone autorizzate, istruite e vincolate alla riservatezza, "
        "secondo il principio di necessità. Mantiene aggiornate autorizzazioni e revoche e assicura una formazione adeguata ai compiti svolti."
    )

    add_heading(doc, "5 Misure tecniche e organizzative", 1)
    doc.add_paragraph(
        "Il Responsabile applica misure adeguate al rischio e le riesamina nel tempo. Le misure effettivamente concordate "
        "sono descritte nell'Allegato A e possono comprendere controllo degli accessi, autenticazione forte, cifratura, backup, "
        "registrazione degli eventi, aggiornamenti, segregazione degli ambienti, procedure di ripristino e gestione degli incidenti."
    )

    add_heading(doc, "6 Sub responsabili", 1)
    doc.add_paragraph(
        "Scegliere una modalità: [autorizzazione specifica preventiva] oppure [autorizzazione generale con obbligo di preavviso "
        "di almeno numero giorni]. Il Responsabile comunica identità, servizio, localizzazione e trattamento affidato. Impone "
        "al sub responsabile obblighi equivalenti e resta responsabile verso il Titolare dell'adempimento degli obblighi affidati."
    )

    add_heading(doc, "7 Assistenza al Titolare", 1)
    doc.add_paragraph(
        "Tenendo conto della natura del trattamento e delle informazioni disponibili, il Responsabile assiste il Titolare nel "
        "riscontro ai diritti degli interessati, nella sicurezza, nella notifica delle violazioni, nelle valutazioni d'impatto e "
        "nelle consultazioni preventive. Le modalità operative e gli eventuali costi ulteriori sono definiti nel contratto principale."
    )

    add_heading(doc, "8 Violazioni dei dati personali", 1)
    doc.add_paragraph(
        "Il Responsabile informa il Titolare senza ingiustificato ritardo e, ove possibile, entro [numero] ore dalla conoscenza "
        "di una violazione. Comunica almeno natura dell'evento, dati e interessati coinvolti, possibili conseguenze, misure adottate "
        "e referente. Conserva le evidenze e collabora agli aggiornamenti, senza effettuare notifiche esterne salvo istruzione o obbligo di legge."
    )

    add_heading(doc, "9 Restituzione e cancellazione", 1)
    doc.add_paragraph(
        "Alla cessazione dei servizi il Responsabile, su scelta del Titolare, restituisce o cancella dati e copie entro [numero] "
        "giorni, salvo obblighi di conservazione. La restituzione avviene nel formato [formato]. Le copie di sicurezza sono eliminate "
        "secondo il ciclo tecnico dichiarato e restano protette fino alla cancellazione."
    )

    add_heading(doc, "10 Informazioni, verifiche e audit", 1)
    doc.add_paragraph(
        "Il Responsabile mette a disposizione le informazioni necessarie a dimostrare il rispetto del presente accordo e "
        "consente verifiche ragionevoli del Titolare o di un soggetto indipendente incaricato. Salvo urgenze o richieste delle "
        "autorità, l'audit è richiesto con [15] giorni di preavviso, durante l'orario di lavoro e tutelando sicurezza e riservatezza di terzi."
    )

    add_heading(doc, "11 Trasferimenti internazionali", 1)
    doc.add_paragraph(
        "Il Responsabile non trasferisce dati fuori dallo Spazio economico europeo senza istruzione documentata e senza una "
        "base conforme alla normativa applicabile. Nell'Allegato B sono indicati Paesi, fornitori, garanzie e misure supplementari eventualmente adottate."
    )

    add_heading(doc, "12 Responsabilità e disposizioni finali", 1)
    doc.add_paragraph(
        "Le responsabilità economiche tra le Parti seguono il contratto principale nei limiti consentiti dalla normativa sulla "
        "protezione dei dati. Modifiche, contatti e istruzioni rilevanti devono risultare per iscritto. Per quanto non disciplinato "
        "si applicano il Regolamento UE 2016 679, la normativa nazionale e le clausole del contratto principale compatibili con questo accordo."
    )

    add_heading(doc, "Allegato A Misure di sicurezza", 1)
    add_table(
        doc,
        ["Area", "Misura applicata o riferimento"],
        [
            ["Accessi e autenticazione", "[Ruoli, MFA, revoca, password]"],
            ["Protezione dati e trasmissioni", "[Cifratura, protocolli, dispositivi]"],
            ["Disponibilità e ripristino", "[Backup, frequenza, test, RTO e RPO]"],
            ["Aggiornamenti e vulnerabilità", "[Patch, scansioni, dipendenze]"],
            ["Log e monitoraggio", "[Eventi registrati e conservazione]"],
            ["Incidenti e continuità", "[Contatti, procedura e test]"],
        ],
        widths=[58, 114],
        font_size=8.2,
    )

    add_heading(doc, "Allegato B Sub responsabili e trasferimenti", 1)
    add_table(
        doc,
        ["Fornitore", "Servizio e dati", "Paese", "Garanzia o base"],
        [["[Nome]", "[Servizio e categorie]", "[Paese]", "[SEE, decisione, SCC o altra base]"]],
        widths=[39, 61, 28, 44],
        font_size=8,
    )

    add_heading(doc, "Sottoscrizione", 1)
    add_table(
        doc,
        ["Per il Titolare", "Per il Responsabile"],
        [["Nome e ruolo: [dati]\nData: [data]\nFirma: ____________________", "Nome e ruolo: [dati]\nData: [data]\nFirma: ____________________"]],
        widths=[86, 86],
        font_size=8.5,
    )
    doc.add_paragraph(
        "Nota d'uso: compilare il modello in base ai trattamenti reali e verificare ruoli, misure, fornitori e trasferimenti con "
        "un professionista privacy. Non utilizzare l'accordo quando VA Digital opera come titolare autonomo e non tratta dati per conto del cliente."
    )
    return save(doc, "11_accordo_trattamento_dati_art28.docx")


def build_change_request():
    doc = setup_document(
        "Richiesta di variazione del progetto",
        "Modulo per descrivere una modifica al perimetro approvato e autorizzarne impatto, costo e tempi prima dell'esecuzione.",
    )
    add_table(
        doc,
        ["Riferimento", "Valore"],
        [
            ["Numero variazione", "[VAR-AAAA-NNN]"],
            ["Cliente e progetto", "[Cliente] - [Progetto]"],
            ["Richiedente", "[Nome, ruolo e contatto]"],
            ["Data e priorità", "[Data] - [Bassa / Media / Alta / Urgente]"],
            ["Contratto o preventivo", "[Riferimento]"],
        ],
        widths=[52, 120],
        font_size=8.8,
    )

    add_heading(doc, "Modifica richiesta", 1)
    add_label_line(doc, "Descrizione", "Indicare cosa deve cambiare e perché")
    add_label_line(doc, "Risultato atteso", "Definire un esito verificabile")
    add_label_line(doc, "Elementi esclusi", "Specificare ciò che non rientra nella richiesta")

    add_heading(doc, "Valutazione di impatto", 1)
    add_table(
        doc,
        ["Area", "Valutazione"],
        [
            ["Funzioni e requisiti", "[Nessun impatto / descrizione]"],
            ["Dati e migrazioni", "[Nessun impatto / descrizione]"],
            ["Integrazioni e infrastruttura", "[Nessun impatto / descrizione]"],
            ["Sicurezza e privacy", "[Nessun impatto / descrizione e verifiche necessarie]"],
            ["Documentazione e formazione", "[Nessun impatto / descrizione]"],
        ],
        widths=[58, 114],
        font_size=8.5,
    )

    doc.add_page_break()
    add_heading(doc, "Effetti economici e temporali", 1)
    add_table(
        doc,
        ["Voce", "Condizione proposta"],
        [
            ["Attività aggiuntive", "[Elenco sintetico]"],
            ["Corrispettivo aggiuntivo", "Euro [importo] oltre imposte e oneri applicabili"],
            ["Pagamento", "[Modalità e scadenza]"],
            ["Nuova data o durata", "[Data / numero giorni aggiuntivi / invariata]"],
            ["Dipendenze del cliente", "[Materiali, accessi, decisioni e relative scadenze]"],
        ],
        widths=[58, 114],
        font_size=8.5,
    )

    add_heading(doc, "Criteri di accettazione", 1)
    add_bullets(doc, ["[Criterio misurabile 1]", "[Criterio misurabile 2]", "[Criterio misurabile 3]"])
    add_heading(doc, "Decisione", 1)
    doc.add_paragraph(
        "[ ] Approvata e autorizzata   [ ] Respinta   [ ] Da rivedere   [ ] Inclusa senza variazione economica"
    )
    doc.add_paragraph(
        "Con l'approvazione, il presente modulo integra il contratto e il preventivo richiamati. VA Digital avvia le attività "
        "solo dopo la sottoscrizione o altra approvazione scritta concordata."
    )
    add_table(
        doc,
        ["Per VA Digital", "Per il Cliente"],
        [["Vincenzo Atturo\nData: [data]\nFirma: ____________________", "Nome e ruolo: [dati]\nData: [data]\nFirma: ____________________"]],
        widths=[86, 86],
        font_size=8.6,
    )
    return save(doc, "12_richiesta_variazione_progetto.docx")


def build_meeting_minutes():
    doc = setup_document(
        "Verbale di riunione",
        "Modello operativo per registrare decisioni, attività, responsabilità e scadenze condivise durante il progetto.",
    )
    add_table(
        doc,
        ["Riferimento", "Valore"],
        [
            ["Cliente e progetto", "[Cliente] - [Progetto]"],
            ["Data, orario e luogo", "[Data] - [Orario] - [Sede o collegamento]"],
            ["Tipo di incontro", "[Avvio / avanzamento / tecnico / collaudo / altro]"],
            ["Redatto da", "[Nome]"],
        ],
        widths=[52, 120],
        font_size=8.8,
    )

    add_heading(doc, "Partecipanti", 1)
    add_table(
        doc,
        ["Nome", "Organizzazione e ruolo", "Presenza"],
        [
            ["[Nome]", "[Azienda e ruolo]", "[Presente / Assente]"],
            ["[Nome]", "[Azienda e ruolo]", "[Presente / Assente]"],
            ["[Nome]", "[Azienda e ruolo]", "[Presente / Assente]"],
        ],
        widths=[52, 82, 38],
        font_size=8.5,
    )

    add_heading(doc, "Ordine del giorno", 1)
    add_bullets(doc, ["[Punto 1]", "[Punto 2]", "[Punto 3]"])

    add_heading(doc, "Sintesi della discussione", 1)
    add_label_line(doc, "Punti emersi", "Riportare fatti, vincoli e alternative valutate")
    add_label_line(doc, "Questioni aperte", "Indicare gli elementi ancora da chiarire")

    add_heading(doc, "Decisioni", 1)
    add_table(
        doc,
        ["N", "Decisione approvata", "Responsabile", "Data effetto"],
        [
            ["1", "[Decisione]", "[Nome]", "[Data]"],
            ["2", "[Decisione]", "[Nome]", "[Data]"],
            ["3", "[Decisione]", "[Nome]", "[Data]"],
        ],
        widths=[12, 94, 40, 26],
        font_size=8.4,
    )

    add_heading(doc, "Attività e scadenze", 1)
    add_table(
        doc,
        ["Attività", "Responsabile", "Scadenza", "Stato"],
        [
            ["[Attività]", "[Nome]", "[Data]", "[Da avviare]"],
            ["[Attività]", "[Nome]", "[Data]", "[Da avviare]"],
            ["[Attività]", "[Nome]", "[Data]", "[Da avviare]"],
            ["[Attività]", "[Nome]", "[Data]", "[Da avviare]"],
        ],
        widths=[78, 42, 26, 26],
        font_size=8.3,
    )

    add_heading(doc, "Prossimo incontro", 1)
    doc.add_paragraph("Data proposta: [data]     Obiettivo: [obiettivo]     Materiali necessari: [elenco]")
    add_heading(doc, "Conferma del verbale", 1)
    doc.add_paragraph(
        "Il verbale viene inviato ai partecipanti il [data]. Eventuali osservazioni devono essere comunicate entro [numero] "
        "giorni lavorativi. In assenza di rilievi, sarà utilizzato come riferimento operativo per le decisioni e le attività indicate."
    )
    add_table(
        doc,
        ["Per VA Digital", "Per il Cliente se richiesta firma"],
        [["Vincenzo Atturo\nData: [data]\nFirma: ____________________", "Nome e ruolo: [dati]\nData: [data]\nFirma: ____________________"]],
        widths=[86, 86],
        font_size=8.5,
    )
    return save(doc, "13_verbale_riunione.docx")


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    paths = [
        build_company_sheet(),
        build_requirements_form(),
        build_proposal(),
        build_kickoff(),
        build_status_report(),
        build_delivery(),
        build_support_sheet(),
        build_letterhead(),
        build_service_contract(),
        build_data_processing_agreement(),
        build_change_request(),
        build_meeting_minutes(),
    ]
    for path in paths:
        print(path)


if __name__ == "__main__":
    main()
